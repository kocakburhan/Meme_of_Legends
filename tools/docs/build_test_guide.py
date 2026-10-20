#!/usr/bin/env python3
"""Rehberi üret: uv run --with python-docx python tools/docs/build_test_guide.py.

Yalnız standart kütüphane ve python-docx gerekir. Kaynaklar belgenin sonunda.
İçindekiler alanını Word'de güncelleyin; sayfa görüntülerini ayrıca render_docx.py
ile üretip kontrol edin. --output ile geçici bir çıktı yolu seçilebilir.
"""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
BASE = "https://mol-game.kocakkburhann.workers.dev"
doc = Document()
page_number = 1
chapters = []
toc_anchor = None


def p(text, style=None):
    return doc.add_paragraph(text, style)


def h(text, level=2):
    doc.add_heading(text, level)


def page(title, level=2):
    global page_number
    doc.add_page_break()
    page_number += 1
    h(title, level)
    if level == 1:
        bookmark = f"chapter_{len(chapters) + 1}"
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(1000 + len(chapters)))
        start.set(qn("w:name"), bookmark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), start.get(qn("w:id")))
        doc.paragraphs[-1]._p.insert(0, start)
        doc.paragraphs[-1]._p.append(end)
        chapters.append((title, bookmark, page_number))


def steps(items, first=1):
    # Her test kendi 1 numarasıyla başlar; gerçek Word numaralandırmasıdır.
    numbering = doc.part.numbering_part.element
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    number_id = max(ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(number_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), "7")  # python-docx'in ondalık List Number tanımı
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), str(first))
    override.append(start)
    num.append(override)
    numbering.append(num)
    for item in items:
        para = p(item, "List Number")
        props = para._p.get_or_add_pPr()
        np = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), "0")
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), str(number_id))
        np.extend([il, ni])
        props.append(np)


def expected(text):
    para = p("")
    para.add_run("Beklenen sonuç: ").bold = True
    para.add_run(text)


def checks(*items):
    for item in items:
        para = p("☐ " + item)
        para.paragraph_format.space_after = Pt(3)


def test(title, items, result, *checklist):
    h(title, 3)
    steps(items)
    expected(result)
    checks(*checklist)


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for cell, title in zip(t.rows[0].cells, headers):
        cell.text = title
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text
    props = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn("w:" + key), value)
        borders.append(edge)
    props.append(borders)
    for index, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                m = OxmlElement("w:" + side)
                m.set(qn("w:w"), "100")
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            cp.append(margins)
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), "E3EAF0" if index == 0 else ("F5F7F9" if index % 2 == 0 else "FFFFFF"))
            cp.append(fill)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.bold = index == 0
    p("").paragraph_format.space_after = Pt(0)


def code(text):
    p(text, "Command")


def field(para, instruction, fallback=""):
    r = para.add_run()._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    ins = OxmlElement("w:instrText")
    ins.set(qn("xml:space"), "preserve")
    ins.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, ins, sep, text, end):
        r.append(element)


def fill_toc():
    """Gerçek TOC alanı ve ilk açılışta da okunabilen önbellek sonucu.

    Sayfa önbelleği açık sayfa sonlarına dayanır ve render ile doğrulanır.
    Word güncellerken TOC'yi gerçek düzen ve başlıklardan yeniden hesaplar.
    """
    opening = toc_anchor.add_run()._r
    for name, value in (("fldChar", "begin"), ("instrText", ' TOC \\o "1-1" \\h \\z \\u '), ("fldChar", "separate")):
        element = OxmlElement("w:" + name)
        if name == "fldChar":
            element.set(qn("w:fldCharType"), value)
        else:
            element.set(qn("xml:space"), "preserve")
            element.text = value
        opening.append(element)
    previous = toc_anchor._p
    for title, bookmark, number in chapters:
        element = OxmlElement("w:p")
        previous.addnext(element)
        para = Paragraph(element, toc_anchor._parent)
        para.paragraph_format.space_after = Pt(13)
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        for key, value in (("val", "right"), ("leader", "dot"), ("pos", "9638")):
            tab.set(qn("w:" + key), value)
        tabs.append(tab)
        para._p.get_or_add_pPr().append(tabs)
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), bookmark)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = title
        run.append(text)
        link.append(run)
        para._p.append(link)
        para.add_run("\t")
        field(para, f" PAGEREF {bookmark} \\h ", str(number))
        previous = element
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    Paragraph(previous, toc_anchor._parent).add_run()._r.append(end)


def setup():
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    normal.element.get_or_add_rPr().append(OxmlElement("w:lang"))
    normal.element.rPr[-1].set(qn("w:val"), "tr-TR")
    for name, size in (("Title", 28), ("Subtitle", 13), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 11.5)):
        s = doc.styles[name]
        s.font.name = "DejaVu Sans"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(12 if name.startswith("Heading") else 0)
        s.paragraph_format.space_after = Pt(7)
        s.paragraph_format.keep_with_next = True
        for border in list(s.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    command = doc.styles.add_style("Command", 1)
    command.font.name = "DejaVu Sans Mono"
    command.font.size = Pt(9)
    command.paragraph_format.space_before = Pt(3)
    command.paragraph_format.space_after = Pt(7)
    command.paragraph_format.keep_together = True
    doc.styles["List Number"].paragraph_format.space_after = Pt(5)
    header = sec.header.paragraphs[0]
    header.text = "MEME OF LEGENDS  •  TEST VE KULLANIM REHBERİ"
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("20 Ekim 2026  •  ").font.size = Pt(8)
    field(footer, " PAGE ", "1")
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    doc.core_properties.title = "Meme of Legends Test ve Kullanım Rehberi"
    doc.core_properties.subject = "Proje ortağı için içerik girişi ve uçtan uca manuel test"
    doc.core_properties.author = "Meme of Legends"


def build():
    global toc_anchor
    setup()
    p("Meme of Legends", "Title")
    p("Test ve Kullanım Rehberi", "Subtitle")
    p("Proje ortağı için adım adım uygulama • 20 Ekim 2026")
    p("Bu rehberle önce oyun içeriğini hazırlayın, sonra admin ve işletme panellerini kontrol edin, en son telefonlarla gerçek bir oyun oynayın. Teknik bilgi gerektiren tek bölüm, kendi bilgisayarınızdan meme yükleme aracının kurulumudur; o bölümde komutlar da tek tek açıklanır.")
    p("Belge bir test sonucunu değil, sizin uygulayacağınız testleri anlatır. Bir adım beklediğiniz gibi çalışmazsa kutuyu boş bırakın ve 9. bölümdeki şablonla bildirin. Canlı sistemde yalnız size ait test işletmesi ve test içerikleri üzerinde işlem yapın.")
    h("Hızlı erişim")
    table(["Ekran", "Adres"], [
        ("Oyuna katılım", BASE),
        ("Admin girişi", BASE + "/admin/giris"),
        ("İşletme girişi", BASE + "/isletme/giris"),
        ("Ortak ekran", BASE + "/ekran"),
    ], [3.4, 13.6])
    p("Admin kullanıcı adı: mol-admin. Şifre bu belgede yok; proje sahibi (owner) ayrıca iletir. Aynı admin hesabı ortak kullanılır.")
    p("Önerilen sıra: 1–3 ile hazırlanın → 4 ile birkaç meme yükleyin → 5–6 ile panelleri deneyin → 7A ile tam oyun oynayın → 7B–F ile özel durumları kontrol edin → 8–10 ile görünümü, hata kaydını ve temizliği tamamlayın.")
    checks("Test tarihi: __________________", "Testi yapan: __________________")

    page("İçindekiler", 2)
    p("Word başlıkları gezinme bölmesinde de görünür. İçindekilerde ana bölümler, sayfalarda ayrıntılı alt başlıklar bulunur.")
    toc_anchor = p("")
    p("Sayfa numaraları görünmüyorsa Word masaüstünde içindekiler alanına sağ tıklayın, Alanı Güncelleştir ve Tüm tabloyu güncelleştir seçin. Bölüm başlıkları Başlık 1, alt başlıklar Başlık 2 ve testler Başlık 3 stilleridir.")

    page("1 Giriş", 1)
    p("Meme of Legends, kafe, bar, şirket veya etkinlik gibi bir işletmenin düzenlediği canlı meme yarışmasıdır. Her tur ekranda kısa bir durum belirir. Örneğin ‘Maaşın yattığı an’ durumuna oyuncular en uygun veya en komik içeriği seçer.")
    p("Gönderiler önce kimin seçtiği açıklanmadan gösterilir. Oyuncular ve seyirciler, uygun oldukları gönderilere 1–5 puan verir. Turun sonunda oylar kilitlenir, gönderenler ve puanlar açıklanır. Lig, eleme veya hibrit modu oyunun kazananını belirler.")
    p("Katılımcıların uygulama indirmesi veya hesap açması gerekmez; tarayıcı ve oda kodu ya da QR yeterlidir. İşletmenin oyun açabilmesi için hesabını admin oluşturur ve oyun hakkı verir. Oyun başladıktan sonra aşamalar kendiliğinden ilerler; yönetici gerektiğinde müdahale eder.")
    h("Roller")
    table(["Rol", "Ne yapar"], [
        ("Platform admin", "İşletme açar, hak verir, tüm içeriği ve raporları yönetir; gerektiğinde oturumu sonlandırır."),
        ("İşletme yöneticisi", "Lobiyi açar, ayarları yapar, konsoldan oyunu ve katılımcıları yönetir."),
        ("Oyuncu", "Hesapsız katılır; meme gönderir ve kurala uygun diğer gönderileri puanlar."),
        ("Seyirci", "Gönderi yapmaz; geç katılım kuralına göre puan verir."),
        ("Ortak ekran", "TV veya projektörde sahneyi gösterir; yönetim veya puanlama yapmaz."),
    ], [4.2, 12.8])
    h("Kısa sözlük")
    table(["Terim", "Anlamı"], [
        ("Oda kodu", "Katılım için kullanılan 6 karakter; ekran eşleştirme kodundan farklıdır."),
        ("Durum ve gönderi", "Durum turun konusu; gönderi oyuncunun o konu için seçtiği meme'dir."),
        ("Tur ve oturum", "Tur seçimden sonuca kadar bir döngü; oturum lobi ile birlikte bir oyundur."),
        ("Lig / eleme / hibrit", "Lig toplam puan; eleme düello galipleri; hibrit önce lig, sonra yarı final ve finaldir."),
        ("Özel içerik", "Oyuncunun yüklediği resim/GIF veya seçtiği YouTube klibi; kütüphaneye eklenmez."),
        ("Ortak ekran ve konsol", "Ortak ekran izleme yüzeyi; konsol işletmenin oyun yönetim ekranıdır."),
    ], [4.2, 12.8])

    page("2 Teste hazırlık", 1)
    test("Cihazları ve kişileri hazırlayın", [
        "Bir bilgisayar, 3–4 telefon ve internet hazırlayın. TV veya ikinci ekran yoksa ortak ekranı ayrı bir tarayıcı penceresinde açın.",
        "Telefonlarda güncel Chrome veya Safari; bilgisayarda güncel Chrome, Edge veya Safari kullanın. İlk turda cihaz adlarını ve tarayıcılarını not alın.",
        "Bilgisayarda admin ve işletme panellerini ayrı sekmelerde açın. Katılımcılar için farklı telefonlar ya da farklı tarayıcı profilleri kullanın. Sekmelere Ada, Bora, Cem ve Deniz gibi adlar verin.",
        "Telefon azsa normal pencere ile gizli pencereyi ayrı kişi olarak kullanabilirsiniz. Aynı tarayıcının birden çok gizli penceresi genellikle aynı kimliği paylaşır; her pencereyi ayrı kişi saymayın. Üçüncü ve dördüncü kişi için farklı profil veya farklı tarayıcı açın.",
        "Katılımcı pencerelerini test boyunca kapatmayın; gizli pencerelerin tamamını kapatmak kimliği kaybettirebilir. Konsol veya ortak ekran, oyuncu yerine geçmez; katılımı ayrıca yapın.",
    ], "Her katılımcı bağımsız bir kimlikle bağlanabilir. Yönetici ve ortak ekran pencereleri ayrıca açık kalır.",
        "3 oyuncu ve 1 seyirci için ayrı kimlik hazır", "Admin şifresi owner'dan alındı", "Ortak ekran için pencere hazır")
    h("Test notlarını tutun")
    p("Her testin kutusunu yalnız beklenen sonucu gördüğünüzde işaretleyin. Yapamadığınız test için ‘Denenmedi’ ve nedenini yazın. İlgili oda kodunu, cihazı ve saati birlikte kaydedin; sonradan raporu bulmak kolaylaşır.")
    p("Testler canlı ortamda hak tüketir. İlk işletmeyi sonlu hakla oluşturun; hak düşmesi ve iade testinden sonra gerekiyorsa sınırsız hak açın. Pasife alma, şifre sıfırlama, oyuncu atma ve sonlandırma işlemlerini yalnız test kayıtlarında deneyin.")
    checks("Test kayıtlarında TEST öneki kullanılacak", "Not dosyası veya kâğıt hazır")

    page("3 İlk içeriği ve test işletmesini hazırlama", 1)
    test("Kategori ve platform durumları", [
        "Admin girişini açın; mol-admin ve owner'ın ilettiği şifreyle giriş yapın. Menüden İçerik, ardından Kategoriler bölümünü açın.",
        "Kategori adı alanına TEST Komik yazıp Ekle'ye basın. Daha sonra gerçek kütüphane için anlaşılır kategori adları oluşturun.",
        "İçerik içinden Platform durumları bölümünü açın. Yeni durum alanına ‘TEST Kahve bitince toplantının başlaması’ yazın; Ekle'ye basın.",
        "Oyuna uygun farklı durumları tek tek ekleyin. Metinler 3–160 karakter olmalı; link veya HTML kullanmayın. İlk deneme için en az 30 aktif platform durumu önerilir; daha zengin bir açılış havuzu için 100 ve üzeri hedeflenebilir.",
        "Genel bakışa dönün. Aktif platform durumu sayısını ve az durum uyarısını kontrol edin. 30'un altında uyarı çıkar; bu sert bir oyun başlatma sınırı değildir.",
    ], "Kategori ve durumlar listede görünür. Durumlar aktif havuza girer. Havuz turlarda yenilenir; tüm durumlar kullanılırsa oyun eski durumları tekrar kullanarak devam eder.",
        "Kategori oluşturuldu", "Aktif durum sayısı kontrol edildi", "Test durumu kaydedildi")
    test("Meme kütüphanesini ve işletmeyi hazırlayın", [
        "4. bölümdeki yerel araçla birkaç resim, GIF ve kısa video ekleyin. Kategori ve açıklaması eksik olanları Hızlı etiketleme ile tamamlayın. İlk oyun için birkaç aktif ve açılabilir meme bulunduğunu kontrol edin.",
        "İşletmeler bölümünde oluşturma formunu açın. İşletme adı olarak ‘TEST Rehber 20 Ekim’, benzersiz kullanıcı adı olarak ‘test-rehber-20ekim’ yazın; ad alınmışsa sonuna farklı bir sayı ekleyin.",
        "Başlangıç hakkını 10 bırakacak şekilde girin. Önerilen şifreyi güvenli bir yere kaydedin ve Oluştur'a basın. Giriş bilgileri özetini kopyalama düğmesiyle alın.",
    ], "Test işletmesi aktiftir. Özet işletme giriş adresini, kullanıcı adını ve geçici şifreyi içerir. İşletme kendi kendine kayıt olmaz.",
        "En az bir resim, GIF ve video hazır", "Test işletmesinin giriş bilgileri kaydedildi")

    page("4 Bilgisayardan meme ekleme", 1)
    p("Bu kurulum yalnız kütüphaneye meme ekleyeceğiniz bilgisayarda gerekir. Oyun oynayan telefonlarda kurulum yapmayın. Terminal, komut yazıp Enter'a bastığınız penceredir. Aşağıdaki komutları satır satır çalıştırın; bir komut bitmeden sonrakine geçmeyin.")
    h("Node.js kurulumunu işletim sisteminize göre yapın")
    p("Resmî indirme sayfası: https://nodejs.org/en/download . Araç Node.js 22 veya üzerini ister. Aşağıdaki yol için 22 serisinin güncel bakım sürümünü seçmek Corepack ile başlangıcı kolaylaştırır. 22.13 ve üzerini kullanın; farklı bir yeni sürümde Corepack ayrıca kurulabilir.")
    table(["Sistem", "Kısa kurulum"], [
        ("Windows", "Sitede Windows'u ve bilgisayarınıza uygun mimariyi seçin; .msi kurucusunu indirin. Kurucuyu açıp varsayılan adımları tamamlayın. Terminali kapatıp yeniden açın."),
        ("macOS", "Sitede macOS'u seçin. Apple Silicon için ARM64, Intel için x64 seçin; .pkg kurucusunu açıp adımları tamamlayın. Terminal'i yeniden açın."),
        ("Linux", "Sitede Linux ve uygun mimariyi seçin; 22 serisinin .tar.xz hazır dosyasını indirin. Dosya yöneticisinde çıkarın. İçinde bin klasörü olan ana klasörü node22 olarak yeniden adlandırıp kişisel klasörünüze taşıyın. Terminalde aşağıdaki Linux komutunu çalıştırın."),
    ], [2.8, 14.2])
    p("Yalnız Linux için, her yeni terminalde önce bu satırı çalıştırın; ardından diğer sistemlerle ortak adımlara geçin. node22 klasörü kişisel klasörünüzde olmalıdır.")
    code('export PATH="$HOME/node22/bin:$PATH"')
    steps(["Terminalde aşağıdaki sürüm kontrolünü çalıştırın. Çıktı v22.13 veya daha yeni bir sürüm olmalıdır."])
    code("node --version")
    steps(["pnpm adlı yardımcıyı etkinleştirin. Corepack, projenin istediği pnpm sürümünü seçer."], first=2)
    code("corepack enable")
    p("‘corepack bulunamadı’ veya imza hatası çıkarsa aşağıdaki komutla Corepack'i kurun/güncelleyin, sonra corepack enable komutunu yeniden çalıştırın. Windows kurucusunun mevcut Corepack bileşeni çakışırsa owner'dan destek alın.")
    code("npm install --global corepack@latest")
    p("Yetki reddi alırsanız kurulumu owner ile tamamlayın. Windows PowerShell komut dosyası çalıştırma hatası verirse aynı işlemi Komut İstemi'nde deneyin.")
    expected("node --version bir sürüm döndürür; Corepack komutu hatasız tamamlanır.")
    checks("Node.js kuruldu", "Corepack etkin")
    p("Kurulum kaynakları: https://nodejs.org/en/download ve https://github.com/nodejs/corepack#manual-installs . pnpm kurulum seçenekleri: https://pnpm.io/installation .")

    page("Depoyu edinme ve aracı hazırlama")
    test("Proje klasörünü bilgisayarınıza alın", [
        "Owner'dan projenin GitHub bağlantısını ve erişim iznini isteyin. Bu rehber depo adresi tahmin etmez; owner'ın verdiği bağlantıyı kullanın.",
        "Git bilmeniz gerekmez: GitHub'da oturum açın, owner'ın belirttiği sürümü açın, Code menüsünden Download ZIP ile indirin. ZIP'i açıp tüm dosyaları normal bir klasöre çıkarın. Komutları ZIP'in içinden çalıştırmayın.",
        "Owner isterse projeyi Git ile sizin için alabilir. Sonraki adımlar iki durumda da aynı proje klasöründe yapılır. Güncelleme gerektiğinde owner'dan yeni sürümü isteyin.",
        "İçinde package.json, pnpm-lock.yaml, apps ve tools bulunan klasörü açın. Burası proje kökü; tools/media alt klasöründe başlamayın.",
        "Windows'ta bu klasörün adres çubuğuna cmd yazıp Enter'a basın. macOS'ta Finder'da klasöre sağ tıklayıp Hizmetler menüsündeki klasörde terminal açma seçeneğini kullanın. Linux'ta dosya yöneticisinde sağ tıklayıp terminal açma seçeneğini seçin. Seçenek yoksa owner ile klasöre geçin.",
    ], "Terminal doğru proje klasöründe açılır; proje dosyaları dosya yöneticisinde görünür.", "ZIP çıkarıldı veya depo alındı", "Terminal proje kökünde")
    h("Bağımlılıkları ve medya yardımcılarını kurun")
    steps(["Aşağıdaki komutu çalıştırın. İlk indirme onayı gelirse onaylayın; internet hızına göre birkaç dakika sürebilir."])
    code("pnpm install")
    steps(["Ardından medya yardımcılarını indirin. Araç ffmpeg/ffprobe ile videoyu işler; yt-dlp ile bağlantıdan indirir. Bunları ayrı ayrı kurmanız genellikle gerekmez."], first=2)
    code("pnpm media setup")
    steps(["Canlı sunucuya giriş yapın. Kullanıcı adı mol-admin; şifre owner'ın ilettiği admin şifresidir. Şifre yazarken karakter görünmemesi normaldir; yazıp Enter'a basın."], first=3)
    code("pnpm media login --server " + BASE)
    expected("Kurulum ve giriş tamamlanır. Tarayıcıda admin'e girmiş olmak bu aracın girişinin yerine geçmez. Şifre kaydedilmez; bilgisayarda oturum bilgisi tutulur.")
    checks("pnpm install tamamlandı", "Medya yardımcıları hazır", "Canlı sunucuda araç girişi başarılı")
    p("Windows'ta yerel medya aracı henüz denenmedi. Kurulum veya işlem hatasını komut ve ekran görüntüsüyle owner'a bildirin. Linux x64 doğrulanmıştır; diğer cihazlarda bu rehberdeki sonuçları ayrıca kontrol edin.")

    page("Tek meme ekleme ve video kesme")
    p("Komutlardaki LINK yerine gerçek bağlantıyı, kategori ve açıklama yerine kendi metninizi yazın. Açılı ayraçlar veya LINK sözcüğü gerçek bağlantının parçası değildir. Boşluk içeren metinleri düz çift tırnak içinde tutun.")
    h("Bağlantıdan veya dosyadan ekleyin")
    steps(["Önce admin panelinde kategoriyi oluşturun. Ardından YouTube, X veya Instagram gönderi bağlantısıyla aşağıdaki kalıbı kullanın."])
    code('pnpm media add "LINK" --category "TEST Komik"\n--description "TEST İlk gün"')
    p("Yukarıdaki iki görsel satırı terminale tek satır olarak yazın. Çalışır örnek aşağıdadır; 5–20. saniyeleri alır.")
    code('pnpm media add "https://youtu.be/jNQXAC9IVRw" --start 0:05 --end 0:20 --category "TEST Komik" --description "TEST İlk gün"')
    steps(["Bilgisayarınızdaki bir resmi eklemek için dosyayı proje klasörüne kopyalayın; aşağıdaki dosya adını kendi dosyanızla değiştirin."], first=2)
    code('pnpm media add "./resim.png" --category "TEST Komik" --description "TEST Şaşkınlık"')
    p("Kategori adı etkin kategorilerle Türkçe büyük/küçük harf duyarsız eşleşir. Kategori ve açıklama birlikte verilmezse dosya etiketlenmemiş kuyruğuna gider; oyunda görünmesi için panelden tamamlayın. Aynı kaynak bağlantısı tekrar eklenirse mükerrer olarak bulunur.")
    h("25 saniye kuralı")
    p("Kütüphane videosu en fazla 25 saniyedir. Uzun videoda başlangıç ve bitişi siz seçmelisiniz. 0:05–0:20 arası 15 saniyedir. YouTube bağlantısındaki başlangıç parametresi kesim sayılmaz; --start ve --end kullanın.")
    code('pnpm media add "./video.mov" --start 0:05 --end 0:20')
    p("Bu son örnek etiket vermediği için kuyruğa gider. --start tek başına videonun kalanını, --end tek başına baştan o ana kadarını alır; sonuç yine 25 saniyeyi aşmamalıdır. Kesim resim/GIF için kullanılmaz.")
    expected("İşlenen dosya yüklenir; etiketleri tamam olan içerik Meme kütüphanesinde aktif görünür ve önizlemesi açılır. Video seçilen aralıkta, en fazla 25 saniye olur.")
    checks("Bağlantıdan yükleme denendi", "Yerel resim yüklendi", "Video kesildi ve panelde açıldı")
    p("Yerel araçtaki X/Instagram desteği, oyuncunun oyun içi özel link alanından farklıdır: oyun içinde özel video bağlantısı yalnız YouTube'dur. Çok medyalı gönderilerde araç ilk öğeyi alır. Resim saydamlığı dönüşümde korunmaz; GIF animasyonu korunur.")

    page("Toplu yükleme ve sorun çözme")
    h("Bağlantı listesi ve deneme işlemi")
    steps(["Not Defteri veya düz metin düzenleyicisinde links.txt oluşturup proje köküne kaydedin. Her satıra tek bağlantı yazın. Dosyanın links.txt.txt olmadığını kontrol edin. Boş satırlar ve # ile başlayan açıklamalar atlanır."])
    code("# Deneme listesi\nhttps://youtu.be/jNQXAC9IVRw\nhttps://x.com/NASA/status/1760811461943722132")
    steps(["Listeyi yükleyin. Bu örnekte etiketler eksik olduğu için başarılı öğeler etiketlenmemiş kuyruğuna gider. 25 saniyeyi aşanlar kesim gerekli hatası verir; onları add ile ayrı kesin."], first=2)
    code('pnpm media import "./links.txt" --report "./rapor.json"')
    p("Klasör de verebilirsiniz: pnpm media import \"./arsiv\" . Alt klasörlere girilmez. Bağlantı listesine kesim seçenekleri yazılmaz; import komutunda --start/--end kullanılamaz.")
    steps(["Yükleme yapmadan denemek için aşağıdaki komutu kullanın. Dosya yine indirilip dönüştürülür; sunucuya yüklenmez ve mükerrer kontrolü yapılmaz. Kategori seçeneği verirseniz sunucuya kategori sorgusu yapılır."], first=3)
    code('pnpm media add "./resim.png" --dry-run')
    p("Raporu metin düzenleyicisinde açın. uploaded yüklendi, duplicate zaten var, dry-run yalnız denendi, error başarısız demektir. error yanında neden vardır. Bir öğe hata alsa da diğerleri işlenir. --report vermezseniz mol-media-report-<zaman>.json oluşturulur. Mevcut raporun üzerine yazılmaz; yeni rapor adı seçin.")
    expected("Rapor her kaynak için sonuç verir. Etiketsiz başarılar panelde Hızlı etiketleme kuyruğuna düşer; dry-run dosya eklemez.")
    checks("links.txt biçimi denendi", "Rapor okundu", "dry-run kütüphaneyi değiştirmedi")
    page("Medya aracı hata mesajları")
    h("İndirme bozulursa")
    p("İlk çare yt-dlp'yi güncelleyip aynı komutu bir kez daha denemektir. Gizli veya giriş isteyen bağlantı güncellemeyle açılmayabilir; herkese açık başka içerik ya da yerel dosya kullanın.")
    code("pnpm media update")
    table(["Gerçek mesaj veya mesaj başlangıcı", "Ne yapmalısınız"], [
        ("Kesim gerekli (süre: … sn)", "add komutunda --start/--end ile en fazla 25 sn seçin."),
        ("İçerik giriş veya platform doğrulaması gerektiriyor.", "Herkese açık başka kaynak kullanın; tarayıcı oturumu aktarımı yoktur."),
        ("İçerik gizli; herkese açık bir içerik kullanın.", "Gizli içeriği araç indiremez."),
        ("İçerik silinmiş veya artık erişilemiyor.", "Kaynağı tarayıcıda kontrol edin; başka içerik seçin."),
        ("Desteklenmeyen site veya gönderi biçimi.", "YouTube, X veya Instagram gönderisi ya da yerel dosya kullanın."),
        ("Ağ hatası; bağlantıyı ve platform erişimini kontrol edin.", "İnterneti ve kaynağın tarayıcıdan açıldığını kontrol edin."),
        ("Admin oturumu geçersiz; pnpm media login çalıştırın.", "Canlı adresle login komutunu yeniden çalıştırın."),
        ("yt-dlp bulunamadı. Önce pnpm media setup çalıştırın.", "setup komutunu yeniden çalıştırın."),
        ("MEDIA_INVALID: Sunucu medya türünü, boyutunu veya süresini reddetti.", "Dosyayı ve kesim süresini kontrol edin; düzelmezse owner'a raporlayın."),
    ], [9, 8])
    expected("Güncelleme sonrası uygun kaynak yeniden işlenebilir. Gizli, silinmiş veya giriş isteyen içerik için başka kaynak gerekir. Hata sürerse komutun tamamını ve mesajı owner'a iletin.")
    checks("Güncelleme komutu bulundu", "Hata mesajının karşılığı anlaşıldı")

    page("5 Admin paneli testleri", 1)
    test("Giriş hatası ve çıkış", [
        "Admin giriş adresinde mol-admin ve bilerek yanlış bir şifreyle bir kez Giriş yap'a basın.",
        "Doğru şifreyle giriş yapın; menüde Genel bakış, İşletmeler, Oturumlar, İçerik ve Sistem bölümlerini bulun.",
        "Çıkış yap'a basın. /admin adresini yeniden açıp giriş istendiğini kontrol edin; sonra tekrar giriş yapın.",
    ], "Yanlış şifre hata verir ve panel açılmaz. Doğru giriş paneli açar; çıkıştan sonra korumalı sayfa girişe yönlendirir.", "Hatalı giriş reddedildi", "Giriş ve çıkış çalıştı")
    test("Genel bakış ve telefon menüsü", [
        "Genel bakışta aktif oturumları, platform durum sayısını, etiketlenmemiş meme sayısını ve son hataları inceleyin. Veri yoksa boş liste açıklamasını okuyun.",
        "Bir aktif test oturumu varsa satırından detayına geçin; geri dönün. Etiketlenmemiş öğe varsa ilgili bağlantıyı açın.",
        "Telefondan admin'e giriş yapın. Menü düğmesiyle aynı bölümlere gidin; form alanlarını ve işlemleri kaydırarak kontrol edin.",
    ], "Sayılar ve bağlantılar ilgili kayıtlara gider. Etiket kuyruğu büyükse 100+ yazabilir. Veri yokken boş açıklama görünmesi normaldir; kırık ekran veya sonsuz yükleme beklenmez.", "Genel bakış bağlantıları açıldı", "Telefonda menü ve formlar kullanılabildi")
    h("Bu bölümde hangi kayıt kullanılacak")
    p("3. bölümdeki TEST işletmesini kullanın. Oyun gereken adımlarda 7A ile bir oyun başlatın. Hesap etkileyen testleri ayrı ayrı uygulayın; her seferinde gerekiyorsa yeniden giriş yapın. Gerçek işletmenin şifresini veya hakkını değiştirmeyin.")

    page("İşletme oluşturma ve imtiyazlar")
    test("Oluşturma ve giriş bilgisini kopyalama", [
        "3. bölümde işletme oluşturmadıysanız şimdi İşletmeler'deki oluşturma formunu açın; İşletme adı, Kullanıcı adı, Başlangıç hakkı ve Önerilen şifreyi doldurun; Oluştur'a basın.",
        "Giriş bilgileri özetini kopyalayın, kendi notunuza yapıştırın. Kullanıcı adı ve geçici şifreyi işletme girişinde deneyin.",
        "İşletmeler listesinde ad veya kullanıcı adıyla arayın; Aktif, Pasif ve Tümü filtrelerini deneyin. Test işletmesinin detayını açın.",
    ], "İşletme bir kez oluşturulur, bilgiler kopyalanır ve doğru işletmeye giriş yapılır. Aynı kullanıcı adıyla ikinci oluşturma reddedilir.", "Oluşturma ve kopyalama başarılı", "Arama ve filtre doğru")
    test("Hak ekleme çıkarma ve sınırsız hak", [
        "Detayda mevcut bakiyeyi not edin. Hak miktarı alanına 2 yazıp Hak ekle'ye basın; bakiyenin 2 arttığını kontrol edin. Ardından Hak çıkar ile aynı 2 hakkı geri alın.",
        "Bakiyeden büyük miktar çıkarmayı deneyin; işlem reddedilmeli ve bakiye negatif olmamalıdır.",
        "Sınırsız hak anahtarını açın; hak ekle/çıkarın kapandığını görün. Bir deneme oyununu başlatıp sayısal hak düşmediğini doğrulayın. Sonraki hak iadesi testi için sınırsızı tekrar kapatın.",
    ], "Haklar doğru değişir; sınırsız oyun bakiye tüketmez. Sonlu hakta lobi açmak tüketmez, Başlat onayı tüketir. 0 hakta yeni lobi açılamaz.", "Hak ekle ve çıkar doğru", "Negatif bakiye engellendi", "Sınırsız hak denendi")
    test("Eşzamanlı oturum ve lobi ömrü", [
        "Eşzamanlı oturum sınırı varsayılan 1 iken işletmede bir lobi açın; ikinci lobi girişiminin engellendiğini kontrol edin. Admin'den sınırı 2 yapın ve ikinci lobiyi deneyin. Test lobilerini iptal edin; sınırı 1'e geri alın.",
        "Lobi ömrünü 1 saat yapın ve yeni bir lobi oluşturun. Admin Oturumlar detayındaki oluşturulma ve Lobi sona erme zamanlarını karşılaştırın. Yaklaşık 1 saat fark olmalıdır; eski lobinin zamanının geriye dönük değişmesini beklemeyin.",
        "Süre dolmasını bekleyerek test edebiliyorsanız lobinin kapandığını kontrol edin; bekleyemiyorsanız bu kısmı Denenmedi yazın. Sonraki lobiler için ömrü 24 saate geri alın.",
    ], "Açık lobi ve oyunlar eşzamanlı sınıra dahildir. Yeni lobinin son zamanı ayara uyar; süre bitince başlatılmamış lobi kapanır.", "Oturum sınırı denendi", "Lobi bitiş zamanı doğru", "Lobi süre dolumu denendi veya not edildi")

    page("İşletme hesabına müdahale")
    test("Pasife alma ve yeniden etkinleştirme", [
        "Test işletmesiyle bir lobi açın; admin detayında Pasife al ve Evet, pasife al'ı seçin.",
        "İşletme penceresini yenileyin; girişin kapandığını ve açık lobinin iptal olduğunu kontrol edin. Pasif hesapla yeni giriş deneyin.",
        "Admin'de Aktifleştir'e basıp yeniden giriş yapın. Ayrıca ayrı bir başlamış test oyununda pasife almayı deneyin; oyuncu pencerelerinden oyunun sürdüğünü kontrol edin. Sonra hesabı yeniden etkinleştirin.",
    ], "Pasifleştirme yönetici girişlerini kapatır, açık lobiyi kapatır ve yeni oturumu engeller. Başlamış oyun sonuna kadar devam eder; admin gerekirse ayrıca sonlandırır.", "Pasif hesap giriş yapamadı", "Lobi kapandı", "Başlamış oyun etkilenmedi")
    test("Şifre sıfırlama", [
        "İşletmeyi iki ayrı tarayıcı veya cihazda açık tutun. Admin'de Şifre sıfırla ve Evet, şifreyi sıfırla'ya basın.",
        "Yeni geçici şifreyi Şifreyi kopyala ile kendi notunuza alın. İşletmenin açık pencerelerini yenileyin; yeniden giriş istendiğini kontrol edin.",
        "Eski şifreyle girişin reddedildiğini ve yeni geçici şifrenin çalıştığını doğrulayın. Bu sırada başlamış oyun varsa oyuncuların akışının sürmesini izleyin.",
    ], "Tüm yönetici girişleri kapanır, oyun devam eder. Yeni şifreyle tekrar girilir.", "Eski girişler kapandı", "Yeni şifre çalıştı")
    test("İşletme adına giriş", [
        "Admin'de aktif TEST işletmesinin detayında Adına giriş'e basın. Tarayıcı yeni sekmeyi engellerse bu site için açılır pencereye izin verin.",
        "Yeni sekmede işletme adını ve paneli kontrol edin; varsa Oyuna dön ile konsolu açın.",
        "Admin sekmesine dönüp panelin açık kaldığını ve Sistem işlem kaydında adına giriş kaydı oluştuğunu kontrol edin.",
    ], "İşletme paneli yeni sekmede açılır; admin girişi sürer. Aynı tarayıcıdaki önceki işletme kimliği değişir. Adına giriş 12 saat geçerlidir ve kayda geçer.", "Doğru işletme açıldı", "Admin oturumu korundu", "İşlem kaydı oluştu")

    page("Oturumlar ve kullanım")
    test("Liste filtreleri ve oturum detayı", [
        "Oturumlar bölümünü açın. İşletme, durum ve tarih filtreleriyle TEST işletmesinin oyunlarını bulun. Tarih alanı tarayıcı diline göre gün/ay sırasını farklı gösterebilir.",
        "Bir aktif oturumu açın; oda kodu, işletme, mod, ayarlar, zamanlar ve katılımcıları kontrol edin. Tamamlanmış bir oturumda turları, oynandığı andaki durum metnini, puanları ve sıralamayı karşılaştırın.",
        "Çıkarma, gizleme veya duraklatma denediğiniz oyunda ilgili kayıtları arayın. Özel içerik kullanılan raporda yalnız tür bilgisi kaldığını kontrol edin; özel dosyanın kalıcı kütüphaneye eklenmesini beklemeyin.",
    ], "Filtreler istenen oturumları getirir; detay oyunla tutarlıdır. Canlı oyunda sonuç henüz oluşmamış olabilir. Yedi gün sonra takma adların anonimleşmesi normaldir.", "Filtreler doğru", "Rapor ve oyun bilgileri eşleşti")
    test("Oturumu sonlandırma ve hak iadesi", [
        "Sonlu haklı test işletmesinde oyun başlatın; başlatmadan önceki ve sonraki bakiyeyi not edin. Admin Oturumlar'da doğru oda kodunun detayını açın.",
        "Oturumu sonlandır ve Evet, sonlandır'a basın. Katılımcı ekranlarının sona erdiğini ve admin detayının sonlandırılmış olduğunu kontrol edin.",
        "Bakiyenin kendiliğinden artmadığını görün. Hakkı iade et ve Evet, iade et'i kullanın; bakiyenin 1 arttığını kontrol edin. Sayfayı yenileyip aynı oturumdan ikinci iadenin mümkün olmadığını kontrol edin.",
        "Sınırsız hakla başlatılmış veya yalnız lobi açılıp iptal edilmiş oturumda iade yapılmadığını kontrol edin; bu oturumlarda gerçek hak harcanmaz.",
    ], "Sonlandırma oyunu kapatır, otomatik hak iade etmez. Gerçekten tüketilmiş hak oturum başına yalnız bir kez iade edilir; müdahaleler kayda geçer.", "Oyun sonlandırıldı", "İade tam 1 hak verdi", "Çift iade engellendi")
    test("Kullanım özeti", [
        "Oturumlar sayfasındaki Kullanım bağlantısını açın. TEST işletmesinin oturum, aktif oturum, harcanan ve iade edilen hak sayılarını inceleyin.",
        "Yukarıdaki başlatma ve iade işlemleriyle karşılaştırın. Sınırsız oyunların harcanan hak sayısını artırmadığını kontrol edin.",
    ], "İşletme kullanımı kayıtlarla uyumludur; oturum sayısı ile harcanan hak sayısı her zaman aynı değildir.", "Kullanım özeti kontrol edildi")

    page("İçerik yönetimi")
    test("Kategoriler", [
        "İçerik → Kategoriler'de TEST adıyla bir kategori Ekle'yin. Düzenle ile adını değiştirip Kaydet'e basın.",
        "Yukarı ve aşağı oklarıyla sırasını değiştirin, sayfayı yenileyip sırayı kontrol edin.",
        "Arşivle'ye basın; Arşivli işaretini görün. Geri al ile geri getirin. Gerçek kategoriler üzerinde deneme yapmayın.",
    ], "Ad ve sıra kalıcıdır. Arşivleme kaydı silmez; geri alınabilir.", "Ekleme ve ad değiştirme başarılı", "Sıralama ve arşiv geri alma başarılı")
    test("Platform durumları", [
        "Platform durumları'nda benzersiz bir TEST metni Ekle'yin. Aynı metni baş/son boşluk ve harf büyüklüğü değiştirerek tekrar eklemeyi deneyin.",
        "Durum metninde arama yapın. Düzenle ve Kaydet ile metni değiştirin.",
        "Arşivle'yi seçin; Arşiv filtresinde bulun. Yeniden etkinleştir ile geri alın ve Aktif filtresinde kontrol edin.",
    ], "Normalize edilmiş aynı metin ikinci kez eklenmez; arşivdeki tekrar da yeni kayıt yaratmaz. Aktif durumlar havuza girer; arşivliler yeni çekimlerden çıkar.", "Tekrar reddedildi", "Arama ve düzenleme çalıştı", "Arşiv ve geri alma çalıştı")
    test("İşletme durumlarını denetleme", [
        "Önce işletmenin Durumlarım bölümünden TEST metni ekleyin. Admin'de İçerik → İşletme durumları yolunu açın; işletme ve metin filtresiyle bulun.",
        "Engelle'ye basın. İşletme panelinde engelli durumunu kontrol edin; metni düzenlese bile engelin kalkmadığını görün.",
        "Admin'den Engeli kaldır'a basın. İşletme durumu aktifse yeniden havuza girebileceğini kontrol edin.",
    ], "Admin işletmenin metnini düzenleyemez, yalnız engelleyebilir. Engel arşiv durumundan ayrıdır; işletmenin düzenlemesi engeli aşmaz.", "Engelleme görüldü", "Engel yalnız admin'den kaldırıldı")

    page("Meme kütüphanesi ve sistem kayıtları")
    test("Arama filtre düzenleme ve arşiv", [
        "İçerik → Meme kütüphanesi'nde TEST açıklamasını arayın; kategori, durum ve sıralama filtrelerini deneyin.",
        "Bir test meme'inin önizlemesini açın. Düzenle ile kategori ve açıklamasını değiştirip Kaydet'e basın. Yeni açıklamayla arayın.",
        "Arşivle'ye basıp Arşiv filtresinde bulun. Geri al ile test kaydını geri alın; ilk tam oyun öncesi aktif içerik kaldığını kontrol edin.",
    ], "Arama açıklama üzerinden çalışır. Etkin kayıt için kategori ve açıklama gerekir. Arşivlenen içerik geçmiş kayıtlardan silinmez; yeni seçimlerde aktif içerik kullanılır.", "Önizleme ve filtreler doğru", "Düzenleme kaydedildi", "Arşivleme başarılı")
    test("Hızlı etiketleme ve kısayollar", [
        "Yerel araçla kategori/açıklama vermeden en az iki TEST dosyası yükleyin. İçerik → Hızlı etiketleme'yi açın; büyük önizlemeyi ve varsa kaynak bağlantısını inceleyin.",
        "Kategori seçin ve açıklama girin. Kaydet ve sıradakine geç'e basın; önceki kaydı aktif kütüphanede bulun.",
        "Başka öğede Ctrl+Enter (Mac'te ⌘+Enter) ile kaydedin. Metin alanı dışında 1–9 ile kategori seçimini ve Esc ile atlamayı deneyin. Metin alanında yazarken bu tuşların yazmayı engellememesi gerekir.",
        "Atla'nın öğeyi kuyrukta bıraktığını; Arşivle'nin kuyruktan çıkardığını kontrol edin. Kuyruk boşsa boş açıklamayı görün.",
    ], "Etiketlenen öğe kütüphaneye girer; sıradaki gelir. Kategori veya açıklama eksikken kaydetme açılmaz. Atlanan öğe silinmez.", "Kaydet ve sıradaki çalıştı", "Üç klavye kısayolu denendi", "Atla ve arşiv ayrımı doğru")
    test("Sistem hata ve işlem kayıtları", [
        "Sistem'i açın; hata kayıtlarını ve admin işlem kayıtlarını ayrı ayrı inceleyin.",
        "Hatalarda kaynak, kod veya oturum; işlemlerde işlem veya hedef filtresiyle test kaydını arayın. Filtreleri temizleyip listeyi tekrar kontrol edin.",
        "Hak verme, adına giriş, sonlandırma ve iade işlemlerinin tarihlerini karşılaştırın. Hata kaydı yoksa boş listeyi kontrol edin; hata oluşturmak için gerçek oyunu bozmayın.",
    ], "İşlemler kayda geçer. Hata olmadığı için listenin boş olması normaldir; şifre ve özel içerik bağlantısı kayıt ekranında bulunmamalıdır.", "İşlem kaydı bulundu", "Hata ekranı ve filtreler açıldı")

    page("6 İşletme paneli testleri", 1)
    test("Giriş şifre önerisi ve şifre değiştirme", [
        "İşletme giriş adresini ayrı bir tarayıcıda açın. TEST işletmesinin kullanıcı adı ve geçici şifresiyle Giriş yap'a basın.",
        "Ana sayfada ‘Güvenlik için şifreni değiştirmen önerilir.’ önerisini kontrol edin. Hesap'a gidin; Mevcut şifre, Yeni şifre ve Yeni şifre (tekrar) alanlarını doldurun. Yeni şifre en az 12 karakter olmalıdır.",
        "Önce yeni şifre tekrarını farklı yazın; hata aldığınızı görün. Sonra eşleştirip Şifreyi değiştir'e basın. Yeni şifreyi kendi güvenli notunuza alın.",
        "Bu cihazda açık kaldığınızı kontrol edin. İşletmenin başka cihazdaki girişini yenileyin; yeniden giriş istenmelidir. Yeni şifreyle tekrar giriş yapın.",
    ], "Şifre değişir; değişikliği yaptığınız cihaz açık kalır, diğer cihazlardaki işletme girişleri kapanır. Hatalı tekrar kaydedilmez.", "Şifre önerisi görüldü", "Hata ve başarılı şifre değişimi denendi", "Diğer cihazın girişi kapandı")
    test("Ana sayfa ve yeni oyun", [
        "Ana sayfada işletme adını ve Oyun hakkı değerini kontrol edin. Yeni oyun düğmesini kullanın; oda kodlu konsol açılır.",
        "Lobi açınca hakkın azalmadığını kontrol edin. Ana sayfaya dönün; aktif oturumdaki Oyuna dön ile aynı odaya gidin.",
        "Lobiyi iptal edin; kalan hakkın değişmediğini görün. 7A'da oyunu gerçekten başlatınca sonlu hakta 1 düşmesini kontrol edin.",
    ], "Yeni lobi konsola gider, Oyuna dön mevcut odayı açar. Hak yalnız oyun başlatılırken tüketilir.", "Lobi ve Oyuna dön doğru", "Hak davranışı doğru")

    page("Durumlarım geçmiş ve hesap")
    test("Durumlarım kütüphanesi", [
        "Durumlarım'da ‘TEST Masaya yeni kahve gelince’ yazıp Ekle'ye basın. 3–160 karakter kullanın; link veya HTML girmeyin.",
        "Aynı metni farklı büyük/küçük harf ve fazla boşlukla yeniden eklemeyi deneyin; tekrar hatası bekleyin. 2 karakterlik metin ve bir bağlantı deneyin; kaydedilmemeli.",
        "Metni arayın; Düzenle ile değiştirin, Kaydet'e basın. Arşivle'yi kullanın, Arşiv filtresinde bulun ve Yeniden etkinleştir ile geri alın.",
        "Admin'in engellediği test metnini kontrol edin. Metni değiştirseniz de engelin sürdüğünü görün; engeli admin'den kaldırın.",
    ], "İşletme kendi durumlarını yönetir. Aynı işletme içindeki tekrar reddedilir; platformdaki aynı metin ayrı sahiplikte olabilir. Aktif ve engelsiz metinler oyun havuzuna girer.", "Ekleme doğrulaması ve tekrar denendi", "Arama düzenleme arşiv geri alma çalıştı", "Admin engeli korundu")
    test("Geçmiş oyunlar ve hesap çıkışı", [
        "7A'dan sonra Geçmiş'i açın. Bitmiş oyunun tarihini, modunu ve katılımcı bilgilerini bulun; oyunun detayını açın.",
        "Tur puanlarını, sonuçları ve sıralamayı oyunda aldığınız notla karşılaştırın. Başka bir işletmenin oturumlarının burada görünmediğini kontrol edin.",
        "Hesap bölümünde işletme bilgilerini kontrol edin. Çıkış yap'a basın; yeniden panel açınca giriş istendiğini görün. Sonraki senaryo için tekrar giriş yapın.",
    ], "Geçmiş yalnız bu işletmenin özetlerini gösterir; tamamlanan oyun detayına açılır. Çıkış sonrası panel korunur.", "Geçmiş ve detay tutarlı", "Hesap ve çıkış çalıştı")

    page("7 Tam oyun senaryoları", 1)
    h("Ortak hazırlık")
    steps([
        "İşletmede yeni lobi açın; konsolda Oyun, Katılımcılar ve Ayarlar bölümlerini bulun. TEST dışındaki açık oturumlara dokunmayın.",
        "Konsol üstündeki Ortak ekran bölümünü açın, Eşleştirme kodu üret'e basın. TV veya ayrı pencerede ortak ekran adresini açın; 6 rakamlı kodu girip Eşleştir'e basın. Ekranı başlat ile sesi etkinleştirin.",
        "Eşleştirme kodu 5 dakika geçerli ve tek kullanımlıktır. Süresi dolarsa yeni kod üretin. Aynı anda tek ortak ekran vardır; yeni eşleşme eskisini düşürür. Bu kod telefonların katıldığı oda kodu değildir.",
        "Telefonlarda QR okutun veya ana canlı adreste 6 karakterli oda kodunu yazın. Farklı takma adlarla üç kişi Oyuncu, bir kişi Seyirci seçsin. Takma adlar 2–20 karakter ve oda içinde benzersiz olmalı.",
        "Oyuncu veya seyirci onayı açıksa konsolda bekleyenleri Onayla ile kabul edin. İlk denemede onayları kapalı tutabilirsiniz. Kadroyu oyuncu ve seyirci listesinde kontrol edin.",
        "Konsolda Katılım linkini kopyala, QR göster ve QR indir seçeneklerini deneyin. Ortak ekranın lobiyi, QR'ı ve doğru kişileri gösterdiğini kontrol edin.",
    ])
    expected("Telefonlarda lobi, konsolda kişiler ve ayarlar, ortak ekranda aynı oda görünür. Ortak ekran yönetim yapamaz. Lobi hak tüketmez.")
    checks("QR ve kodla katılım denendi", "Ortak ekran eşleşti", "3 oyuncu ve 1 seyirci hazır")
    h("Başlatmadan önce kısa kural özeti")
    p("Lig ve eleme en az 2, hibrit en az 4 oyuncuyla başlar. Seyirci zorunlu değildir. Başlayınca yeni oyuncu alınmaz; yeni gelen yalnız seyirci olur. Aktif oyuncu sayısı 2'nin altına düşerse oyun hemen biter.")
    p("Gönderi puanı geçerli 1–5 oyların ortalamasıdır; ekranda 2 ondalık görünür. Oy gelmezse veya gönderi yapılmazsa 0 olur. Kendi gönderinize puan veremezsiniz; her geçerli oyun ağırlığı eşittir.")

    page("Süre puanlama ve beraberlik başvuru notları")
    table(["Konu", "Beklenecek davranış"], [
        ("Seçim", "30–240 sn; varsayılan 60. Herkes onaylarsa erken biter. Onaylanmasa bile süre sonunda son geçerli seçim gider; geçerli seçim yoksa 0."),
        ("Gösterim", "Resim/GIF 12+3 = 15 sn. Video klip süresi, en az 12 sn, ardından 5 sn son karar: toplam 17–30 sn. Son kararda oy değişebilir."),
        ("Final gözden geçirme", "15–120 sn; varsayılan 45. Eksik oy tamamlanır, oy değişir; gönderen gizli kalır. Tüm bağlı uygun kişiler Bitirdim derse erken biter."),
        ("Sonuç", "Oylar kilitlenir, gönderenler ve ortalamalar açıklanır. Açıklama bittikten sonra 15 sn geri sayımla ilerler; yönetici bekletebilir veya hemen geçebilir."),
        ("Düello", "Rakipler birbirine oy vermez; tarafsız diğer kişiler verir. Odada tarafsız katılımcı yoksa iki rakip birbirine puan verir. 2 oyunculu lig de düellodur."),
        ("Geç seyirci", "Gönderim gösterimi başladıktan sonra gelen kişi o turu, gözden geçirme dahil, puanlayamaz. Sonraki tur puanlar. Seçimde gelen o tur puanlayabilir."),
    ], [4, 13])
    h("Beraberliği nasıl anlarsınız")
    p("Eleme düellosu, hibrit maç, hibritte ilk 4'e giriş sınırı ve lig birinciliği çözülür. Diğer sıralardaki eşitlik paylaşılır. Hibritin ilk 4'ü içindeki sıra eşitliği kurayla belirlenir.")
    p("Hibrit maç toplamı eşitse önce lig aşaması toplamı yüksek olan kazanır. Gerekirse yeni durumla tek uzatma turu gelir: seçim 30 saniyedir, yalnız berabere kalanlar gönderi yapar. Uzatma da eşitse kura belirler; yönetici kazanan seçmez. Aynı turdaki birden fazla beraberlik tek uzatmada çözülür. Uzatma puanları toplamlara veya jüri ödüllerine eklenmez.")
    h("Jüri ödülleri")
    p("En sert jüri ve En cömert jüri için kişinin oyun boyunca en az 5 puan vermesi gerekir. Koşulu sağlayan en az iki kişi yoksa ödüller görünmez; aynı kişi iki ödülü alamaz. Tek turluk 3 oyunculu denemede ödüllerin çıkmaması hata değildir.")

    page("Senaryo A Lig ile ilk tam oyun")
    p("Kadroyu 3 oyuncu ve 1 seyirci yapın. Ayarlarda Lig, 1 tur ve özel içerik kapalı seçin; süreleri ilk denemede varsayılan bırakın. Beraberlik istemiyorsanız puanlayıcılar belirgin biçimde farklı puanlar versin.")
    table(["Kim", "Ne yapar"], [
        ("Yönetici", "Ayarları kontrol eder, Başlat → Evet, başlat ile hak onayını verir; sonra akışı izler."),
        ("Ada Bora Cem", "Kütüphaneden meme seçer; kendi gönderisi dışında puan verir."),
        ("Deniz seyirci", "Gönderi yapmadan üç içeriği puanlar; gözden geçirmeyi tamamlar."),
        ("Ortak ekran", "Durum, anonim gönderiler, sonuç ve oyun sonunu gösterir."),
    ], [4, 13])
    steps([
        "Başlat onayından önce ve sonra kalan hakkı not edin; sonlu hakta 1 azalmalı. Her telefonda aynı durum ve seçim geri sayımı görünmeli.",
        "Oyuncular Kütüphane'de kategori ve aramayı denesin, bir meme'i önizlesin ve Seçimimi onayla'ya bassın. Son oyuncu onaylamadan biri seçimini değiştirsin; seçim açık olduğu sürece değişebilmeli.",
        "Son oyuncu onaylayınca gösterimin erken başladığını görün. Gönderiler tüm ekranlarda aynı sırada, gönderen adı gizli olarak gösterilmeli. Herkes uygun gönderilerde 1–5 düğmelerinden puan versin; Deniz bir puanını son karar sırasında değiştirsin.",
        "Canlı oy sayacını kontrol edin; ortalamalar henüz görünmemeli. Ada kendi gönderisinde puan verememeli. Deniz bir gönderiyi puanlamadan bıraksın.",
        "Final gözden geçirmede Deniz kaçırdığı gönderiyi açıp puanlasın; Bora bir puanını değiştirsin. Herkes Bitirdim'e bassın. Aşama bitince eski oyların değiştirilemediğini kontrol edin.",
        "Sonuçta gönderenler, geçerli oy sayısı ve iki ondalıklı ortalamalar açıklansın. Bir gönderiye verilen bilinen oyların ortalamasını notunuzla karşılaştırın. Tek tur bittiğinde final sıralaması ve şampiyon görünmeli; birincilik eşitse uzatma/kura bekleyin.",
        "Yönetici Rövanş'a bassın. Oda kodu aynı kalmalı, kişiler odadan çıkmadan yeni lobiye geçmeli. Oyuncular yeniden oyuncu olma isteği göndermeli. Yeni oyunu başlatırsanız yeniden 1 hak tüketilir; yalnız lobi açılırsa tüketilmez.",
    ])
    expected("Seçim → anonim gösterim → gözden geçirme → sonuç → oyun sonu tamamlanır. Seyirci puan verir; ligde kimse elenmez. Rövanş aynı odayı yeni lobiye taşır.")
    checks("Tam akış bütün ekranlarda tamamlandı", "Oy değiştirme ve kilitleme doğru", "Puan ortalaması doğru", "Oyun sonu ve rövanş doğru")

    page("Senaryo B Özel içerik ve YouTube")
    p("Yeni bir Lig oyunu açın; 3 oyuncu ve 1 seyirci, 1 tur, özel içerik açık, seçim 240 sn olsun. Hazır bir PNG/JPG ve herkese açık gömülebilir YouTube bağlantısı bulundurun. Son oyuncu, ret ve yeniden seçim bitmeden onaylamasın.")
    table(["Kim", "Ne yapar"], [
        ("Ada", "Yükle sekmesinden resim gönderir; reddedilince kütüphaneden yeniden seçer."),
        ("Bora", "YouTube linkini doğrular, en fazla 25 sn kırpar ve seçer."),
        ("Cem", "Kütüphaneden seçim yapar; denemeler tamamlanana kadar onayı bekletir."),
        ("Yönetici ve Deniz", "Yönetici reddeder; Deniz anonim gösterimi ve puanlamayı kontrol eder."),
    ], [4, 13])
    steps([
        "Oyunu başlatın. Ada Yükle'yi açıp dosya seçsin; yükleme yüzdesini ve önizlemeyi görsün. Bu içeriği seç ve ardından Seçimimi onayla ile göndersin. Dosyanın yalnız yüklenmesi seçim onayı sayılmaz.",
        "Yönetici Özel içerik (canlı) listesinde Ada'nın adını ve önizlemesini görsün; Reddet → Evet, reddet'i kullansın. Ada kabul edilmedi bildirimini görüp süre içinde Kütüphane'den yeni seçim yapsın.",
        "Bora YouTube sekmesinde link girip Linki doğrula'ya bassın. Video süresi yüklenince Başlangıç ve Bitiş çizgileriyle örneğin 5–20 saniyeyi seçsin; önizlesin, klibi seçip Seçimimi onayla'ya bassın. 25 saniyeden uzun aralığın kabul edilmediğini kontrol edin.",
        "Vakit varsa erişilemeyen/gizli bir bağlantı ve 10 MB üstü resim deneyin; anlaşılır hata bekleyin. Resim/GIF için JPG, PNG, WebP ve GIF desteklenir; oyuncu video dosyası yükleyemez.",
        "Herkes onaylayınca Ada'nın yeni seçiminin ve Bora'nın YouTube klibinin anonim gösterildiğini kontrol edin. YouTube seçili aralıktan başlamalı, bitişi geçmemeli. Ortak ekran varsa ses oradan gelsin; telefonlar sessiz başlayabilir, isteyen Sesi aç ile açabilsin.",
        "Oyunu bitirin. İşletme geçmişi ve admin raporunda özel içerik türünü kontrol edin. Özel dosyalar ve YouTube ayrıntıları oturum bitince silinir; kalıcı meme kütüphanesine eklenmez.",
    ])
    expected("Yönetici bakmasa da oyun ilerler. Reddedilen içerik geçersiz olur; seçim süresi varsa oyuncu yeniden seçebilir, yoksa o tur 0 alır. YouTube tüm yüzeylerde oynar; cihazlar arasında kare kare eşzamanlılık beklenmez.")
    checks("Resim yükleme ve seçim onayı başarılı", "Ret bildirimi ve yeniden seçim çalıştı", "YouTube doğrulama kırpma ve gösterim başarılı", "Özel içerik kütüphaneye eklenmedi")

    page("Senaryo C Dört oyuncuyla eleme")
    p("Dört oyuncu için dört bağımsız tarayıcı kimliği gerekir. İsterseniz ek seyirci kullanın. Yeni lobide Eleme seçin; tur sayısını sistem belirler.")
    table(["Kim", "Ne yapar"], [
        ("Yönetici", "4 oyuncuyu kontrol eder ve oyunu başlatır; sonuç seçmez."),
        ("4 oyuncu", "İlk tur seçim yapar; kendi düellosunu puanlamaz, diğer düelloyu puanlar."),
        ("Elenen 2 kişi", "Seyirciye dönüşür; final düellosuna tarafsız oy verir."),
        ("Ortak ekran", "Eşleşmeleri, anonim gösterimi, düello sonuçlarını ve şampiyonu gösterir."),
    ], [4, 13])
    steps([
        "Başlatın; ilk turdaki iki rastgele eşleşmeyi not edin. Her oyuncu bir meme seçip onaylasın.",
        "Gösterimde oyuncuların kendi gönderisine ve düello rakibine puan veremediğini, diğer eşleşmedeki gönderilere verebildiğini kontrol edin. Eşleşmeler açıklansa da gösterim sırasında gönderen gizli kalmalı.",
        "Gözden geçirmeyi bitirin; her düelloda yüksek ortalamalı gönderinin kazandığını görün. Eşitlik olursa tek uzatma ve gerekirse kura bekleyin.",
        "Kaybedenlerin seyirci olduğunu ve sonraki turda gönderi yapamadığını kontrol edin. İki galip finalde yeni seçim yapsın; önceki turun puanı taşınmamalı.",
        "Finalde elenen seyirciler puan versin. Şampiyon 1., final kaybedeni 2. görünmeli. Aynı turda elenen diğerleri o tur puanına göre sıralanmalı; eşitse sıra paylaşılmalı.",
    ])
    expected("Dört oyuncu iki düellodan iki finaliste iner. Elenenler oyunda seyirci kalır. Tarafsız kişi bulunduğu için finalistler birbirine oy veremez.")
    checks("Eşleşmeler ve oy yetkisi doğru", "Elenme seyirciye geçirdi", "Final ve sıralama doğru")
    h("Ek deneme")
    p("Vaktiniz varsa 3 oyuncuyla ayrı bir Eleme açın. Bir kişi bay geçmeli, gönderi yapmadan sonraki tura çıkmalı ve diğer düelloyu puanlayabilmeli. İki oyunculu, hiç seyircisiz ayrı denemede ise tarafsız kişi olmadığı için rakiplerin birbirini puanlayabildiğini doğrulayın.")
    checks("Bay denendi veya Denenmedi notu yazıldı", "Tarafsız kişi olmayan düello denendi veya not edildi")

    page("Senaryo D Hibrit ligden finale")
    p("Hibrit en az 4 oyuncu ister; ilk 4 dışındakilerin elenmesini de görmek için mümkünse 6 veya daha fazla oyuncu kullanın. Yalnız 4 kimlik varsa oynanabilir; ilk 4 dışı elenme kontrolünü Denenmedi yazın. Hibrit normalde 5+3+3 = 11 turdur; uzatma eklenebilir.")
    table(["Kim", "Ne yapar"], [
        ("Yönetici", "Hibrit ayarını seçer; oyuncu sayısı ve aşama geçişlerini izler."),
        ("Tüm oyuncular", "İlk 5 tur lig oynar; her tur uygun gönderilere puan verir."),
        ("İlk 4", "1–4 ve 2–3 eşleşmelerinde 3 tur paralel yarı final oynar."),
        ("Diğerleri ve finalistler", "Elenenler puanlar; iki galip 3 tur final oynar."),
    ], [4, 13])
    steps([
        "Hibrit lobisini başlatın. İlk 5 tur herkes seçim ve puanlama yapsın; lig toplamlarının biriktiğini kontrol edin.",
        "Beşinci turun sonunda ilk 4'ün kaldığını, diğerlerinin seyirciye dönüştüğünü görün. 4./5. sınırı eşitse uzatma/kura; ilk 4 içindeki sıra eşitse kura bekleyin.",
        "Yarı finalde 1–4 ve 2–3 eşleşmelerini kontrol edin. Dört oyuncu aynı durum için eşzamanlı seçim yapmalı. Oyuncu kendi rakibinin gönderisini puanlayamamalı; diğer eşleşmedekileri puanlayabilmeli.",
        "Üç turun sonunda eşleşme toplamlarını karşılaştırın. Lig puanı maç toplamına eklenmez; maç eşitliğinde ilk ayırıcı olarak kullanılır.",
        "İki galibin 3 tur final oynadığını, kaybedenlerin seyirci olarak puan verdiğini görün. Maç toplamı eşitse önce lig toplamı, gerekirse uzatma/kura uygulanmalı.",
        "Final galibi 1., kaybedeni 2. olmalı. Yarı final kaybedenlerinden lig puanı yüksek olan 3., diğeri 4. olur; lig puanı eşitse paylaşırlar. 5. ve sonrası lig sırasıdır.",
    ])
    expected("Lig, iki paralel yarı final ve final otomatik ilerler. Her maçın puanı kendi üç turundan oluşur. İlk 4 belli olmadan oyuncu sayısı 4'ün altına düşerse yarı finaller iptal edilir ve lig tablosu final olur.")
    checks("5 turluk lig tamamlandı", "İlk 4 ve elenme doğru", "Paralel yarı final ve oy yetkisi doğru", "3 turluk final ve sıralama doğru")
    p("İlk 4 belirlendikten sonra yarı finalist ayrılırsa rakibi hükmen finale çıkar; finalist ayrılırsa diğeri şampiyon olur. Bu ayrılma testlerini ana akışı bitirdikten sonra ayrı oyunda deneyin.")

    page("Senaryo E Konsol araçları ve süre yönetimi")
    p("Yeni bir Lig oyununu en az 4 oyuncu ve 1 seyirciyle açın; en az 3 tur ve uzun seçim süresi seçin. İşlemleri sırayla yapın. Oyuncuyu seyirciye alma ve atma işlemlerini sona bırakın; yarışan sayısı 2'nin altına düşerse oyun biter.")
    table(["Kim", "Ne yapar"], [
        ("Yönetici", "Önizleme, süre, durum ve duraklatma araçlarını kullanır."),
        ("Bir oyuncu", "Seçimini erken yapar; acil durum değişiminde sıfırlandığını kontrol eder."),
        ("Diğerleri", "Son kişi onaylamayı bekletir; tüm yüzeylerin davranışını izler."),
    ], [4, 13])
    steps([
        "Lobide İlk turun durumu (önizleme) bölümünden Durumu değiştir'i kullanın. Tur başına tek değiştirme hakkı vardır; bu ilk turda acil değiştirme hakkını da tüketir.",
        "Sonraki turda önizlemede değiştirme yapmayın. Seçimin ilk 15 saniyesinde oyuncu bir seçim yapsın; yönetici Durumu değiştir'i basılı tutarak uygulasın. Yeni durum gelmeli, seçimler sıfırlanmalı ve süre baştan başlamalı. İkinci değiştirme yapılamamalı.",
        "Seçimde +15 sn, +30 sn ve +60 sn'yi sırayla deneyin; kalan sürenin arttığını kontrol edin. Seçim bitince bu araçların kullanılamadığını görün; sayaç kısaltma yoktur.",
        "Seçim aşamasında Duraklat'a basın; telefonlarda seçim ve oyun hareketlerinin durduğunu görün. Devam ile seçim kalan süreden sürmeli.",
        "Bir gönderi gösterilirken yeniden Duraklat ve Devam kullanın; o gönderinin gösterimi baştan başlamalı. Final gözden geçirmede duraklatıp devam ederseniz kalan süreden sürmeli.",
        "Sonuç ekranında bir kez Beklet, ardından Devam kullanın. Başka sonuçta Hemen geç'i deneyin. Sonraki turun durum önizlemesinde değiştirme yaparsanız oyuncular henüz görmeden değişmeli.",
    ])
    expected("İşlemler tüm ekranlarda aynı aşamayı korur. Önizleme ve acil değiştirme toplamda tur başına bir haktır; acil değiştirme yalnız ilk 15 saniyededir. Duraklatmada seçim ve oy dondurulur.")
    checks("Önizleme ve acil değiştirme doğru", "Üç süre ekleme seçeneği çalıştı", "Duraklat ve devam doğru", "Sonuç bekletme ve geçiş doğru")

    page("Senaryo E İçerik gizleme ve katılımcı araçları")
    table(["Kim", "Ne yapar"], [
        ("Yönetici", "Gösterimde gizler; Katılımcılar'dan ad, rol ve çıkarma işlemlerini yapar."),
        ("Test oyuncusu ve seyirci", "Rol veya ad değişimini ve geri giriş davranışını kontrol eder."),
        ("Diğer cihazlar", "Gizlenen içeriğin kalktığını ve oyunun devamını izler."),
    ], [4, 13])
    steps([
        "Gösterim sırasında Hemen gizle'ye kısa dokunun; kazara gizleme olmamalı. Sonra 1 saniye basılı tutun; içerik tüm ekranlardan kalkmalı. Özel içerik deneyecekseniz yeni lobide özel içeriği açık tutun.",
        "Özel içerik gizlendiğinde oyuncunun tur puanı 0 olmalı. Sistem meme'i gizlendiğinde ceza yerine telafi uygulanır: lig/hibritte aynı aşamanın diğer turlarının ortalaması aşama sonunda hesaplanır; başka tur yoksa 0. Elemede gizlenen sistem meme'i düellosu uzatmaya gider.",
        "Katılımcılar'da bir test oyuncusuna Adı sıfırla uygulayın. Geçici ad ve yeni ad istemi görünmeli; oyuncu yeni benzersiz ad seçsin. Rolü ve puanı korunmalı.",
        "Oyunculardan birine Seyirciye al uygulayın. Odada kalmalı, yarışmayı bırakmalı ve uygun gönderileri puanlayabilmeli. Geçmiş turlar değişmemeli; o turun kilitlenmemiş gönderisi geçersiz olur.",
        "Test seyircisinde At → Evet, at kullanın. Aynı tarayıcıdan tekrar girememeli; o turdaki kilitlenmemiş oyları silinir. Eski kilitli tur sonuçları değişmemeli.",
        "Konsolda Ayarlar → Oyunu bitir → Evet, bitir'i seçin. Oyun duraklatılmışsa önce Devam'a basın. Oyun erken bitmeli; sonlu hak otomatik iade edilmemeli. Admin detayında gizleme, çıkarma, duraklatma ve erken bitiş kayıtlarını kontrol edin.",
    ])
    expected("Gizleme herkese yansır. Seyirciye alma, atma ve ad sıfırlama ayrı sonuçlar verir. Oyunu bitirme onayla çalışır. Rapor müdahaleleri saklar.")
    checks("Kısa dokunma gizlemedi ve 1 sn basma gizledi", "Özel içerik cezası veya sistem telafisi denendi", "Ad sıfırlama rol ve puanı korudu", "Seyirciye alma ve atma doğru", "Erken bitiş raporda")

    page("Senaryo F Bağlantı ve cihaz sorunları")
    p("En az 2 turluk Lig açın; 3 oyuncu ve 1 seyirciyle başlayın. Geç katılımı denemek için ayrıca temiz bir tarayıcı profili bulundurun. İnternet kesme denemesini yalnız bir telefonda yapın.")
    table(["Kim", "Ne yapar"], [
        ("Ada", "Telefonu yeniler; sonra başka cihazdan aynı adla devam ister."),
        ("Yönetici", "Cihaz devrini onaylar; ortak ekranı yeniden eşleştirir."),
        ("Yeni seyirci", "Gösterim başladıktan sonra katılır; bu ve sonraki turdaki oy yetkisini kontrol eder."),
        ("Bir telefon ve ortak ekran", "Kısa bağlantı kaybından sonra güncel aşamaya döner."),
    ], [4, 13])
    steps([
        "Ada seçim yaptıktan sonra telefon sayfasını yenilesin. Aynı takma ad, rol ve o tur seçimiyle dönmeli; yeni kişi oluşmamalı. Gösterimde yenilemeyi de deneyin; güncel gönderiye dönmeli.",
        "Farklı cihaz/profilde aynı oda ve Ada adıyla giriş deneyin. Adın kullanımda olduğunu ve devam onayı isteğini görün. İsteği gönderin; yönetici Katılımcılar'daki cihaz devri isteğini onaylasın. Yeni cihaz rol, puan ve seçimi devralmalı; eski cihaz yetkisini kaybetmeli.",
        "Bir turun ilk gönderisi gösterilmeye başladıktan sonra yeni seyirci katılsın. O turu izlesin ama final gözden geçirme dahil puanlayamasın. Sonraki turda puanlayabildiğini kontrol edin.",
        "Ortak ekran penceresini kısa süre çevrimdışı bırakıp yeniden açın. Oyun telefonlarda sürmeli; ekran güncel sahneye dönmeli. Eşleşme kaybolduysa konsoldan yeni kodla tekrar bağlayın. İkinci ekranı eşleştirirseniz eskisinin düştüğünü kontrol edin.",
        "Bir telefonda Wi-Fi ve mobil veriyi 5–10 saniye kapatıp geri açın. Oyun diğer cihazlarda sürmeli; kopan cihaz güncel aşamaya dönmeli. Zamanı geçen bir eylemin reddedilmesi normaldir; geçmiş aşamaya oy yazılmamalı.",
        "Yönetici sekmesini kapatıp yeniden Oyuna dön ile açın. Oyunun yönetici bağlantısını beklemeden sürdüğünü kontrol edin. Sunucuyu bu test için yeniden başlatmayın.",
    ])
    expected("Kısa kopma kişiyi otomatik atmaz; seçim kaçırılırsa o tur 0 alabilir. Devretme yönetici onayı ister. Ortak ekran ve yönetici kopması akışı durdurmaz. Tüm katılımcılar 10 dakika kopuk kalırsa oyun otomatik sonlanabilir.")
    checks("Yenileme aynı kişiye döndü", "Cihaz devri ve eski cihazın düşmesi doğru", "Geç seyirci sonraki turda puanladı", "Ortak ekran geri döndü", "Telefon ve yönetici tekrar bağlandı")

    page("8 Erişilebilirlik ve görünüm kontrolleri", 1)
    test("Telefon masaüstü ve yazı büyütme", [
        "Telefon ve masaüstünde katılım, seçim, puanlama, konsol, işletme ve admin panellerini açın. Uzun durum metni ve uzun takma adla taşma olup olmadığını inceleyin.",
        "Telefonu dikey ve yatay çevirin. Gönderi, 1–5 puanlar, süre ve ana düğmeler erişilebilir kalmalı; istemeden başka düğmeye dokunmamalısınız.",
        "Masaüstünde tarayıcı yakınlaştırmasını %200 yapın; telefonda yazı boyutunu büyütmeyi deneyin. Metinler ve hata açıklamaları kesilmemeli, düğmeler ulaşılamaz hale gelmemeli.",
        "Ortak ekranı TV veya 16:9 pencerede kontrol edin. QR, oda kodu ve durum okunmalı. Sesi kapatıp bir tur izleyin; oyunun hangi aşamada olduğu anlaşılmalı.",
    ], "Temel işlemler her düzende yapılabilir. Yakınlaştırmada metin ve kontroller üst üste binmez; bilgi yalnız renkle anlatılmaz.", "Telefon dikey ve yatay düzeni uygun", "Masaüstü ve %200 büyütme kullanılabilir", "Ortak ekran ve sessiz akış anlaşılır")
    test("Klavye ile gezinme", [
        "Masaüstünde fareyi bırakın; Tab ve Shift+Tab ile form ve düğmelerde dolaşın. Hangi öğede olduğunuz görünür bir odakla anlaşılmalı.",
        "Enter veya boşlukla uygun düğmeyi çalıştırın. Seçim sekmelerinde ok tuşlarını deneyin; paneldeki etiketleme kısayollarını da kontrol edin.",
        "Bir onay penceresi açın; odağın pencere içinde kaldığını ve kapatınca önceki düğmeye döndüğünü kontrol edin. Hemen gizle gibi basılı tutulan düğmeyi klavyeyle deneyin.",
    ], "Klavye ile işlem tamamlanır, odak kaybolmaz ve onaylar anlaşılırdır. Sorunu ekran adı ve basılan tuş sırasıyla raporlayın.", "Odak ve klavye akışı kullanılabilir", "Onay penceresi ve geri dönüş doğru")

    page("9 Hata bildirme ve acil durum", 1)
    p("Bir test başarısızsa aynı işlemi üst üste yapmadan önce saati ve oda kodunu not edin. Mümkünse kısa ekran videosu alın. Şifreyi veya giriş bilgisi özetini görüntüye dahil etmeyin. Owner'a aşağıdaki bilgileri birlikte iletin.")
    table(["Alan", "Doldurun"], [
        ("Tarih ve saat", "____________________  Türkiye saati"),
        ("Cihaz ve tarayıcı", "Örn. iPhone / Safari; bilgisayar / Chrome; sürüm biliniyorsa ekleyin"),
        ("Ekran ve bölüm", "Örn. oyuncu / YouTube veya admin / Oturumlar"),
        ("Oda kodu ve işletme", "____________________"),
        ("Tekrar adımları", "1. …  2. …  3. …"),
        ("Beklenen", "Ne görmeniz gerekiyordu?"),
        ("Gerçekleşen", "Ne oldu; ekranda hangi hata yazdı?"),
        ("Kanıt ve sıklık", "Ekran görüntüsü/video; her seferinde mi, bir kez mi?"),
    ], [4.8, 12.2])
    test("Oyun takılırsa uygulanacak yol", [
        "Önce yalnız bir cihazın mı yoksa tüm oyunun mu takıldığını kontrol edin. Bir cihazda medya açılmıyorsa bağlantıyı ve yeniden denemeyi kontrol edin; diğer cihazlar ilerliyorsa oyunu hemen kapatmayın.",
        "Konsola erişebiliyorsanız Duraklat ile mevcut durumu koruyun; hata notunu alın. Devam ile akışa dönmeyi deneyin.",
        "Oyun kurtarılamıyorsa admin panelinde Oturumlar'ı açın, doğru oda kodunu bulun. Oturumu sonlandır → Evet, sonlandır ile yalnız bu test oturumunu kapatın.",
        "Gerçekten hak harcanmışsa ve iade gerekiyorsa Hakkı iade et'i kullanın. Yeni oyun açmadan önce önceki oturumun kapandığını kontrol edin.",
    ], "Sorun kayda alınır ve gerektiğinde doğru oturum kapatılır. Sonlandırma otomatik iade değildir; admin iadesi ayrı işlemdir.", "Bildirim şablonu dolduruldu veya örnek hazırlandı", "Acil sonlandırma yolu bulundu")
    p("‘Bu cihazda yüklenemedi’ tek cihazdaki medya sorununu gösterebilir; oyun durmaz. Hiç geçerli oy gelmeyen gönderi 0 puan alır. Hata raporunda hangi cihazların medyayı görebildiğini de belirtin.")

    page("10 Test sonrası temizlik", 1)
    test("Test kayıtlarını düzenleyin", [
        "Açık test oyunlarını bitirin veya admin'den sonlandırın. Başlatılmamış test lobilerini Lobiyi iptal et ile kapatın; gerçek etkinlik oturumuna dokunmayın.",
        "İşletmeler'de oluşturduğunuz TEST hesaplarını Pasife al ile kapatın. Şifre sıfırlamak temizlik yerine geçmez.",
        "Test platform durumlarını ve test meme'lerini arşivleyin. İşletmenin test durumlarını kendi panelinden arşivleyin. Gerçek içerikleri aktif bırakın.",
        "Mevcut deneme kayıtlarını da inceleyin: ‘Yük testi (orchestrator)’ işletmesi (kullanıcı adı yuk-testi), ‘Canlı test …’ işletmeleri ve ‘Toplantı yine uzayınca’ platform durumu. Owner'ın artık kullanmadığını doğruladığı test kayıtlarını pasife alın veya arşivleyin; gerçek içerikle karıştırmayın.",
        "Test sonuçlarını owner'a iletin: geçenler, başarısızlar ve denenmeyenler ayrı görülsün. Rapor dosyalarını ve hata görüntülerini kendi test notunuzla saklayın.",
    ], "Aktif test lobisi kalmaz; test işletmeleri pasif ve deneme içerikleri arşivlidir. Geçmiş sonuçlar silinmez. Özel dosyalar ve YouTube ayrıntıları oyun bitince temizlenir; takma adlar 7 gün sonra anonimleşir.",
        "Açık test oturumları kapatıldı", "Test hesapları pasife alındı", "Test durumları ve meme'leri arşivlendi", "Geçen başarısız ve denenmeyen testler bildirildi")
    h("Uygulama sırasında belirsiz kalanları yazın")
    p("Windows medya aracı bu proje kapsamında henüz denenmedi. YouTube'un gömme, bölge veya yaş kısıtları cihazdan cihaza değişebilir; gerçek video oynatımını kendi telefon ve ortak ekranınızda kontrol edin. Lobi ömrünü bekleyerek, 7 günlük anonimleşmeyi ve 6+ oyunculu hibriti deneyemediyseniz test raporunda açıkça belirtin.")
    h("Rehberin dayandığı kaynaklar")
    p("Ürün ve kurallar: docs/PRODUCT_SPEC.md tamamı. Kararlar ve canlı ortam: docs/PROJECT_LOG.md son kayıtları; apps/server/README.md Canlı ortam bölümü. Medya komutları: tools/media/README.md ve gerçek hata metinleri için tools/media/src. Admin ve özel içerik sınırları: docs/contracts/admin-api.md ile docs/contracts/custom-content.md. Adresler: apps/web/src/App.tsx. Düğme adları: apps/web/src/pages, console, screens ve components altındaki arayüz kaynakları.")
    p("Bu belge 20 Ekim 2026 tarihli proje kaynaklarına göre hazırlanmıştır. Arayüz veya kurallar değişirse tools/docs/build_test_guide.py güncellenerek yeniden üretilmelidir.")
    fill_toc()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "docs/Meme_of_Legends_Test_Rehberi.docx")
    args = parser.parse_args()
    build()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    all_text = [para.text for para in doc.paragraphs]
    all_text += [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
    print(f"Oluşturuldu: {args.output}")
    print(f"Yaklaşık kelime sayısı: {len(' '.join(all_text).split())}")
